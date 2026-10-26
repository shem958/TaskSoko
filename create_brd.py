import sys
import subprocess

def install(package):
    subprocess.check_call([sys.executable, "-m", "pip", "install", package])

try:
    import docx
except ImportError:
    install('python-docx')
    import docx

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_heading(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    return heading

def create_brd():
    doc = Document()
    
    # Title
    title = doc.add_heading('Business Requirements Document (BRD)', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph('Taskoko - Local Services Marketplace')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_page_break()
    
    # 1. Executive Summary
    add_heading(doc, '1. Executive Summary')
    doc.add_paragraph(
        "Taskoko is Nairobi's premier local services marketplace designed to connect "
        "customers with verified service providers such as plumbers, electricians, and tutors. "
        "The platform enables instant booking, secure in-app payments via M-Pesa, real-time "
        "chat, and user reviews. It bridges the gap between skilled workers and customers "
        "needing immediate or scheduled services, ensuring safety through provider verification (KYC)."
    )
    
    # 2. Project Objectives
    add_heading(doc, '2. Project Objectives')
    doc.add_paragraph("The primary objectives of the Taskoko platform include:", style='List Bullet')
    doc.add_paragraph("Provide a centralized, easy-to-use digital marketplace for local services.", style='List Bullet')
    doc.add_paragraph("Ensure safety and quality through a rigorous provider KYC and verification process.", style='List Bullet')
    doc.add_paragraph("Facilitate seamless cashless transactions via M-Pesa integration.", style='List Bullet')
    doc.add_paragraph("Build a transparent ecosystem through customer reviews and ratings.", style='List Bullet')
    doc.add_paragraph("Offer cross-platform accessibility via Web (PWA) and Mobile apps (Android/iOS).", style='List Bullet')
    
    # 3. Scope of the System
    add_heading(doc, '3. Scope of the System')
    doc.add_paragraph(
        "The system comprises a centralized backend, a web application, and a mobile application. "
        "Key capabilities include User Management, Service Discovery, Booking Management, "
        "In-App Messaging, Payments, and Reviews."
    )
    
    # 4. Functional Requirements
    add_heading(doc, '4. Functional Requirements')
    
    add_heading(doc, '4.1 User Management', level=2)
    doc.add_paragraph("Role-based access: Customers, Providers, and Admins.", style='List Bullet')
    doc.add_paragraph("Customers can register, browse, book services, pay, and leave reviews.", style='List Bullet')
    doc.add_paragraph("Providers must complete KYC (ID number, business registration) and require Admin approval before listing services.", style='List Bullet')
    
    add_heading(doc, '4.2 Service Discovery and Listing', level=2)
    doc.add_paragraph("Providers can list services with descriptions, prices, and categories.", style='List Bullet')
    doc.add_paragraph("Services include geolocation details (latitude, longitude) for map-based discovery.", style='List Bullet')
    doc.add_paragraph("Customers can search and filter services by category, location, and rating.", style='List Bullet')
    
    add_heading(doc, '4.3 Booking System', level=2)
    doc.add_paragraph("Customers can schedule services for a specific date and time.", style='List Bullet')
    doc.add_paragraph("Bookings progress through statuses: PENDING, ACCEPTED, COMPLETED, CANCELLED.", style='List Bullet')
    doc.add_paragraph("Providers can accept or reject pending bookings.", style='List Bullet')
    
    add_heading(doc, '4.4 Payments (M-Pesa Integration)', level=2)
    doc.add_paragraph("Automated M-Pesa STK Push integration via Daraja API.", style='List Bullet')
    doc.add_paragraph("Payment status tracking (PENDING, COMPLETED, FAILED).", style='List Bullet')
    doc.add_paragraph("Transaction recording with Merchant Request IDs and Receipt Numbers.", style='List Bullet')
    
    add_heading(doc, '4.5 In-App Messaging', level=2)
    doc.add_paragraph("Real-time chat functionality linked to specific bookings.", style='List Bullet')
    doc.add_paragraph("Customers and Providers can communicate directly once a booking is initiated.", style='List Bullet')
    
    add_heading(doc, '4.6 Ratings and Reviews', level=2)
    doc.add_paragraph("Customers can rate providers (1 to 5 stars) and leave text feedback after service completion.", style='List Bullet')
    doc.add_paragraph("Service ratings are dynamically updated based on average customer scores.", style='List Bullet')
    
    # 5. Non-Functional Requirements
    add_heading(doc, '5. Non-Functional Requirements')
    doc.add_paragraph("Performance: The system should handle concurrent users and process M-Pesa callbacks within seconds.", style='List Bullet')
    doc.add_paragraph("Security: Secure authentication, encrypted data transmission, and safe storage of KYC details.", style='List Bullet')
    doc.add_paragraph("Scalability: Dockerized infrastructure (PostgreSQL + Redis) enabling easy scaling of backend services.", style='List Bullet')
    doc.add_paragraph("Usability: Responsive web frontend (Next.js PWA) and native mobile apps (React Native Expo) for an optimal user experience.", style='List Bullet')
    
    # 6. Technology Stack
    add_heading(doc, '6. Technology Stack')
    doc.add_paragraph("Backend: Python, Django REST API.", style='List Bullet')
    doc.add_paragraph("Frontend: Next.js (Web / PWA).", style='List Bullet')
    doc.add_paragraph("Mobile App: React Native (Expo).", style='List Bullet')
    doc.add_paragraph("Database & Caching: PostgreSQL, Redis.", style='List Bullet')
    doc.add_paragraph("Infrastructure: Docker & Docker Compose.", style='List Bullet')
    doc.add_paragraph("Third-Party Integrations: M-Pesa Daraja API (Payments), Google Maps API (Location).", style='List Bullet')
    
    # 7. Core Workflows
    add_heading(doc, '7. Core Workflows')
    doc.add_paragraph("Customer Workflow: Register → Search Service → Book → Pay via M-Pesa → Chat with Provider → Receive Service → Review.", style='List Number')
    doc.add_paragraph("Provider Workflow: Register → Submit KYC → Admin Approval → List Service → Accept Booking → Deliver Service → Get Paid.", style='List Number')

    doc.save('Taskoko_BRD.docx')
    print("BRD successfully generated as 'Taskoko_BRD.docx'")

if __name__ == '__main__':
    create_brd()
